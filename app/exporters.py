"""ส่งออกเอกสารเป็น Word (.docx) และรูป JPEG"""
from __future__ import annotations

import io
import threading
from pathlib import Path
from typing import Callable

import pymupdf


def _render_jpeg(page: pymupdf.Page, dpi: int, quality: int) -> bytes:
    s = dpi / 72.0
    pix = page.get_pixmap(matrix=pymupdf.Matrix(s, s), alpha=False)
    return pix.tobytes("jpeg", jpg_quality=quality)


def export_jpeg(doc: pymupdf.Document, target: str, pages: list[int], dpi: int = 200, quality: int = 90,
                progress: Callable[[int, int], None] | None = None) -> list[str]:
    """หน้าเดียว → ไฟล์ตามชื่อที่เลือก, หลายหน้า → ชื่อ_หน้า001.jpg, ชื่อ_หน้า002.jpg ... ในโฟลเดอร์เดียวกัน"""
    t = Path(target)
    out: list[str] = []
    width = max(3, len(str(doc.page_count)))
    for n, i in enumerate(pages):
        if progress:
            progress(n, len(pages))
        path = t if len(pages) == 1 else t.with_name(f"{t.stem}_หน้า{str(i + 1).zfill(width)}{t.suffix or '.jpg'}")
        path.write_bytes(_render_jpeg(doc[i], dpi, quality))
        out.append(str(path))
    return out


def export_docx_images(doc: pymupdf.Document, target: str, pages: list[int], dpi: int = 150,
                       progress: Callable[[int, int], None] | None = None) -> None:
    """Word แบบเหมือนต้นฉบับ: 1 หน้า PDF = 1 หน้า Word เป็นรูปเต็มหน้า (แก้ข้อความไม่ได้แต่หน้าตาตรงเป๊ะ)"""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.shared import Emu

    d = Document()
    first = True
    for n, i in enumerate(pages):
        if progress:
            progress(n, len(pages))
        page = doc[i]
        r = page.rect
        sec = d.sections[0] if first else d.add_section(WD_SECTION.NEW_PAGE)
        first = False
        sec.page_width, sec.page_height = Emu(int(r.width * 12700)), Emu(int(r.height * 12700))
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Emu(0)
        sec.header_distance = sec.footer_distance = Emu(0)
        para = d.add_paragraph()
        para.paragraph_format.space_before = para.paragraph_format.space_after = Emu(0)
        para.paragraph_format.line_spacing = 1.0
        para.add_run().add_picture(io.BytesIO(_render_jpeg(page, dpi, 85)), width=Emu(int(r.width * 12700)))
    d.save(target)


def export_docx_editable(doc: pymupdf.Document, target: str, pages: list[int]) -> None:
    """Word แบบแก้ไขข้อความได้ (แปลงโครงสร้าง/ตาราง/รูปด้วย pdf2docx)"""
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        from pdf2docx import Converter
    import logging
    logging.getLogger().setLevel(logging.WARNING)      # ปิดข้อความ INFO ของ pdf2docx
    tmp = doc.tobytes(garbage=0, deflate=False)
    cv = Converter(stream=tmp)
    try:
        cv.convert(target, pages=pages)
    finally:
        cv.close()


def run_in_thread(fn: Callable[[], None], pump: Callable[[], None]) -> None:
    """รันงานหนักในเธรดแยก แล้ววนเรียก pump() (เช่น processEvents) จนเสร็จ — หน้าต่างจะไม่ค้างระหว่างแปลง"""
    err: list[BaseException] = []

    def work() -> None:
        try:
            fn()
        except BaseException as e:  # noqa: BLE001
            err.append(e)

    th = threading.Thread(target=work, daemon=True)
    th.start()
    while th.is_alive():
        pump()
        th.join(0.05)
    if err:
        raise err[0]
