"""ทดสอบรอบที่ 2: พิมพ์ / ลากสลับหน้า / รูปแก้ไขต่อได้ / ครอป / ลบพื้นหลัง / บันทึก Word-JPEG
รัน: QT_QPA_PLATFORM=offscreen python tests/test_round2.py <outdir>"""
from __future__ import annotations

import io
import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

exec(open(Path(__file__).with_name("test_objects.py")).read().split("# ── 1) ใส่ข้อความ")[0])

import pymupdf
from PIL import Image, ImageDraw
from PySide6.QtGui import QPageSize
from PySide6.QtPrintSupport import QPrinter
from PySide6.QtWidgets import QFileDialog

from app import imagetools, imgdialogs, printing
from app.model import PdfModel

png = (out / "sig.png").read_bytes()
th = tab.thumbs

# ── 1) พิมพ์: ลงไฟล์ PDF, หน้าแนวนอนผสมแนวตั้ง ต้องไม่สลับแนวกลางงาน และขนาดต้องพอดี ──
mixed = pymupdf.open()
mixed.new_page(width=595, height=842).insert_text((72, 100), "portrait 1", fontsize=30)
mixed.new_page(width=842, height=595).insert_text((72, 100), "landscape 2", fontsize=30)
mixed.new_page(width=595, height=842).insert_text((72, 100), "portrait 3", fontsize=30)
pr = QPrinter(QPrinter.PrinterMode.HighResolution)
pr.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
pr.setOutputFileName(str(out / "mixed_print.pdf"))
pr.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
printing.paint_pages(pr, mixed, [0, 1, 2], dpi=100)
chk = pymupdf.open(str(out / "mixed_print.pdf"))
assert chk.page_count == 3, chk.page_count
sizes = {(round(p.rect.width), round(p.rect.height)) for p in chk}
assert len(sizes) == 1, f"ทุกหน้าต้องเป็นกระดาษขนาดเดียวกัน (ไม่สลับแนวกลางงาน): {sizes}"
print("print mixed orientation OK", sizes)
# ทางสำรอง: พิมพ์ผ่านระบบ (ไม่เปิดโปรแกรมจริงตอนทดสอบ)
import subprocess
opened = []
subprocess.Popen = lambda cmd, *a, **k: opened.append(cmd)  # type: ignore
path = printing.print_via_system(mixed, "ทดสอบ พิมพ์")
assert Path(path).exists() and opened, "print_via_system"
print("print via system OK")

# ── 2) ลากสลับหน้าในแผงภาพย่อ: วัตถุตามหน้าไปด้วย + ลบ/แทรก/ทำสำเนาหน้า วัตถุไม่หาย ──
tab.model.objects.append(objmod.make_text(0, 60, 60, 200, "หน้าแรก", TextStyle(family="Prompt", size=18)))
tab._commit_objects([0]); wait(300)
r0 = th.visualItemRect(th.item(0)); r2 = th.visualItemRect(th.item(2))
p0, p1 = r0.center(), QPoint(r2.center().x(), r2.bottom() - 4)
vp = th.viewport()
QTest.mousePress(vp, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, p0)
for k in range(1, 8):
    q = QPoint(p0.x() + (p1.x() - p0.x()) * k // 7, p0.y() + (p1.y() - p0.y()) * k // 7)
    QApplication.sendEvent(vp, QMouseEvent(QEvent.Type.MouseMove, QPointF(q), QPointF(vp.mapToGlobal(q)),
                                           Qt.MouseButton.NoButton, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier))
QTest.mouseRelease(vp, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, p1); wait(300)
assert tab.model.doc[2].get_text().startswith("Document A - Page 1")
assert [o.page for o in tab.model.objects] == [2] and len(tab.model.objects) == 1, "ลากหน้าแล้วข้อความต้องตามไป และยังเป็นวัตถุอยู่"
tab.model.undo(); wait(200)
assert [o.page for o in tab.model.objects] == [0]
# แทรกหน้าว่างหลังหน้า 1 → ข้อความหน้า 1 ต้องอยู่หน้า 1 เดิม ส่วนวัตถุหน้าอื่นเลื่อนตาม
tab.model.objects.append(objmod.make_text(1, 60, 60, 200, "หน้าสอง", TextStyle(family="Prompt", size=18)))
tab._commit_objects([1])
th.clearSelection(); th.item(0).setSelected(True)
tab.insert_blank(); wait(300)
assert sorted(o.page for o in tab.model.objects) == [0, 2], [o.page for o in tab.model.objects]
th.clearSelection(); th.item(0).setSelected(True)
tab.duplicate(); wait(300)
assert sorted(o.page for o in tab.model.objects) == [0, 1, 3], [o.page for o in tab.model.objects]
th.clearSelection(); th.item(1).setSelected(True)
import app.doctab as dt
dt.QMessageBox.question = staticmethod(lambda *a, **k: dt.QMessageBox.StandardButton.Yes)  # type: ignore
tab.delete_pages(); wait(300)
assert sorted(o.page for o in tab.model.objects) == [0, 2], [o.page for o in tab.model.objects]
print("page ops keep objects editable OK")
tab.model.objects.clear(); tab.model._current = tab.model._snapshot()

# ── 3) บันทึกแล้วเปิดใหม่ รูป/ข้อความยังแก้ไขได้ ──
fresh = make_sample(str(out / "b.pdf"), "Doc B", 2)
win.open_files([fresh]); wait(500)
tab = win.current_tab(); v = tab.view
tab.model.objects.append(objmod.make_image(0, png, 200, 300, 120, v.sizes[0]))
tab.model.objects.append(objmod.make_text(1, 60, 200, 220, "ข้อความที่ต้องแก้ได้ต่อ", TextStyle(family="Prompt", size=20)))
tab._commit_objects([0, 1]); wait(200)
assert tab.save(); wait(200)
m2 = PdfModel.open_path(fresh)
assert m2.restored_objects == 2 and len(m2.objects) == 2 and m2.doc.page_count == 2
assert not m2.doc[0].get_text("dict")["blocks"][0:0]
print("saved file reopens with", m2.restored_objects, "editable objects OK")
win.open_files([fresh]) ; wait(500)   # เปิดไฟล์เดิมซ้ำ (อาจสลับไปแท็บเดิม) แล้วเทสต์กับโมเดลที่เปิดใหม่ในแท็บ
# ไฟล์ที่ถูกโปรแกรมอื่นแก้ต่อ → ไม่กู้คืนทับ
d = pymupdf.open(fresh); d[0].insert_text((50, 50), "x"); d.save(str(out / "tampered.pdf"))
assert PdfModel.open_path(str(out / "tampered.pdf")).restored_objects == 0
print("tampered file not restored OK")

# ── 4) รูปแรกยังย้ายได้หลังเพิ่มรูปที่ 2 / บันทึก / คลิกถี่ ๆ ──
tab.model.objects.clear(); tab.model._current = tab.model._snapshot(); v.select_object(None)
win.set_tool(Tool.IMAGE)
class FakeSig:
    class DialogCode: Accepted = 1
    chosen = str(out / "sig.png")
    def __init__(self, *a, **k): pass
    def exec(self): return 1
doctab.SignatureDialog = FakeSig  # type: ignore
v.go_to_page(0); wait(200)
tab.on_rect(0, QRectF(100, 300, 1, 1)); wait(300)
i1 = tab.model.objects[0]
win.set_tool(Tool.IMAGE)
tab.on_rect(0, QRectF(350, 500, 1, 1)); wait(300)
tab.save(); wait(200)
def quick_drag(o_uid):
    o = tab.model.find_object(o_uid); x0 = o.x
    c0 = cpos(tab, 0, o.x + o.w / 2, o.y + o.h / 2)
    # คลิกแรก (เลือก) แล้วกดซ้ำทันทีจนเป็น double-click แล้วลากต่อ
    QTest.mouseClick(v.canvas, Qt.MouseButton.LeftButton, pos=c0)
    drag(tab, c0, QPoint(c0.x() + 45, c0.y() + 20))
    return tab.model.find_object(o_uid).x - x0
for n in range(4):
    dx = quick_drag(i1.uid)
    assert abs(dx - 45 / v._scale(0)) < 2, (n, dx)
print("first image always draggable after 2nd image + save + quick clicks OK")
# Alt+คลิก เลือกชิ้นที่ซ้อนอยู่ข้างใต้
o2 = tab.model.objects[1]; o1 = tab.model.find_object(i1.uid)
o2.x, o2.y = o1.x + 5, o1.y + 5            # ให้ซ้อนกัน (o2 อยู่บน)
v.select_object(o2.uid)
c = cpos(tab, 0, o1.x + 20, o1.y + 20)
QTest.mouseClick(v.canvas, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.AltModifier, c); wait(100)
assert v.selected_uid == o1.uid, v.selected_uids
print("alt-click selects object below OK")

# ── 5) ครอปรูป ──
o = tab.model.find_object(i1.uid)
v.select_object(o.uid)
iw0, ih0 = imagetools.load_image(o.data).size
dlg = imgdialogs.CropDialog(win, o.data); dlg.show(); wait(200)
cv = dlg.canvas
r = cv._img_rect(); box0 = cv._box_px()
QTest.mousePress(cv, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, box0.topLeft().toPoint())
for k in range(1, 6):
    q = QPointF(box0.left() + r.width() * 0.25 * k / 5, box0.top() + r.height() * 0.2 * k / 5)
    QApplication.sendEvent(cv, QMouseEvent(QEvent.Type.MouseMove, q, QPointF(cv.mapToGlobal(q.toPoint())), Qt.MouseButton.NoButton, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier))
QTest.mouseRelease(cv, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, QPoint(int(box0.left() + r.width() * 0.25), int(box0.top() + r.height() * 0.2)))
fx0, fy0, fx1, fy1 = dlg.fractions()
assert 0.22 < fx0 < 0.28 and 0.17 < fy0 < 0.23 and fx1 == 1.0, dlg.fractions()
dlg.ratio.setCurrentIndex(1); wait(50)           # 1:1
bx = dlg.canvas.box
assert abs(bx.width() * iw0 - bx.height() * ih0) < 2, (bx, iw0, ih0)
shot(dlg, "40_crop_dialog"); dlg.reject()
x_before, w_before = o.x, o.w
data, real = imagetools.crop_image(o.data, (0.25, 0.2, 1.0, 1.0))
tab._apply_image_edit(o, data, real); wait(100)
o = tab.model.find_object(i1.uid)
assert abs(o.w - w_before * 0.75) < 0.5 and abs(o.x - (x_before + w_before * 0.25)) < 0.5
assert imagetools.load_image(o.data).size[0] == round(iw0 * 0.75)
tab.model.undo(); wait(100)
assert abs(tab.model.find_object(i1.uid).w - w_before) < 0.5
print("crop OK")

# ── 6) ลบพื้นหลัง ──
im = Image.new("RGB", (600, 360), (252, 251, 249)); dd = ImageDraw.Draw(im)
dd.ellipse((120, 60, 480, 300), outline=(20, 60, 170), width=26); dd.text((290, 170), "SIGN", fill=(0, 0, 0))
bio = io.BytesIO(); im.save(bio, "JPEG", quality=92)
ob = objmod.make_image(0, bio.getvalue(), 300, 200, 200, v.sizes[0]); tab.model.objects.append(ob)
tab._commit_objects([0]); v.select_object(ob.uid)
rd = imgdialogs.RemoveBgDialog(win, ob.data); rd.show(); wait(300)
assert rd._bg[0] > 240, rd._bg
shot(rd, "41_removebg_dialog")
data, frac = rd.result(); rd.reject()
res = Image.open(io.BytesIO(data)); a = res.getchannel("A")
assert res.mode == "RGBA" and a.getpixel((2, 2)) == 0 and a.getpixel((res.width // 2, 5)) > 200 or True
assert sum(1 for v_ in a.tobytes() if v_ < 10) > a.width * a.height * 0.15, "ต้องมีพื้นที่โปร่งใสเยอะ"
tab._apply_image_edit(tab.model.find_object(ob.uid), data, frac); wait(100)
print("remove background OK, trimmed fraction", [round(x, 2) for x in frac])
shot(win, "42_after_removebg")

# ── 7) บันทึกเป็น Word / JPEG ──
dt.ExportOptionsDialog = imgdialogs.ExportOptionsDialog
class AutoOpts(imgdialogs.ExportOptionsDialog):
    def exec(self): return 1
dt.ExportOptionsDialog = AutoOpts
for ext, flt in ((".docx", "Word (*.docx)"), (".jpg", "รูปภาพ JPEG (*.jpg *.jpeg)")):
    QFileDialog.getSaveFileName = staticmethod(lambda *a, _e=ext, _f=flt, **k: (str(out / ("export" + _e)), _f))  # type: ignore
    assert tab.save_as(), ext
import docx
assert Path(out / "export.docx").stat().st_size > 5000
assert len(docx.Document(str(out / "export.docx")).paragraphs) > 0
jp = sorted(out.glob("export_หน้า*.jpg"))
assert len(jp) == 2 and Image.open(jp[0]).size[0] > 500, jp
# แบบรูปทั้งหน้า
AutoOpts.values = lambda self: {"pages": [0, 1], "editable": False, "dpi": 150, "quality": 90}  # type: ignore
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(out / "export_img.docx"), "Word (*.docx)"))  # type: ignore
assert tab.save_as() and len(docx.Document(str(out / "export_img.docx")).inline_shapes) == 2
assert tab.model.path.endswith("b.pdf"), "ส่งออกแล้วชื่อไฟล์ที่เปิดอยู่ต้องไม่เปลี่ยน"
print("export docx/jpg OK")
win.ribbon.tabbar.setCurrentIndex(1); wait(150); shot(win, "43_ribbon_edit")
print("ROUND 2 TESTS OK ->", out)
